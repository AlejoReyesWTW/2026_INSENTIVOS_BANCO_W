"""Fase Novedades: cruce de reemplazos de subgerente en Ventas.

Lee el insumo "Formato novedades subgerente oficina para seguros" (hoja Hoja1,
encabezados en fila 3, registros desde fila 4) y, por cada registro, marca en
la pestaña Ventas las filas donde R == J y la fecha de la transacción (F)
cae dentro del rango [G, H].

LOGICA DE COLUMNAS (validada con datos reales del cliente):
    - J/K (cols 10/11) = QUIEN SE FUE (la persona que ya figura en la columna
      R de Ventas como subgerente). Con su cédula J se BUSCA en R las filas
      de Ventas que le corresponden.
    - A/B (cols 1/2)   = EL REEMPLAZO. Para las filas de Ventas que caen
      dentro del rango [G, H], se ESCRIBE en R (cédula) y S (nombre) los
      datos del reemplazo.
    - G/H (cols 7/8)   = INICIO / FIN del permiso (dd/mm/aaaa).

Columnas de Ventas:
    F (6) = FECHA_TRANSACCION    R (18) = CC_NOVEDAD_SUBGERENTE
    S (19) = NOMBRE_NOVEDAD_SUBGERENTE
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from enum import Enum
from pathlib import Path
from typing import Any

import openpyxl

from src.infrastructure.excel_writer import ExcelWriter

# Columnas del insumo (1-based).
_COL_CEDULA_REEMPLAZO = 1
_COL_NOMBRE_REEMPLAZO = 2
_COL_INICIO = 7
_COL_FIN = 8
_COL_CEDULA_QUIEN_SE_FUE = 10
_COL_NOMBRE_QUIEN_SE_FUE = 11

# Columnas de Ventas (1-based).
_COL_VENTAS_FECHA = 6
_COL_VENTAS_R = 18
_COL_VENTAS_S = 19

_FILA_ENCABEZADO = 3

_FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d")


class EstadoNovedad(str, Enum):
    """Estado del cruce para una persona de novedades."""

    COMPLETO = "COMPLETO"
    PARCIAL = "PARCIAL"
    SIN_DATOS = "SIN DATOS"


@dataclass(frozen=True)
class RegistroNovedad:
    """Un registro del archivo de novedades (un reemplazo)."""

    cedula_quien_se_fue: object
    nombre_quien_se_fue: str
    inicio: date
    fin: date
    cedula_reemplazo: object
    nombre_reemplazo: str


@dataclass(frozen=True)
class ResultadoNovedad:
    """Auditoría del cruce de una persona de novedades."""

    nombre_quien_se_fue: str
    cedula_quien_se_fue: object
    nombre_reemplazo: str
    cedula_reemplazo: object
    encontrados: int
    cambiados: int
    intervalo: str
    estado: EstadoNovedad


def leer_novedades(ruta: Path | str) -> list[RegistroNovedad]:
    """Lee los registros de novedades (desde fila 4).

    Args:
        ruta: ruta al archivo de novedades (.xlsx).

    Returns:
        Lista de registros válidos. Las filas incompletas (sin cédula de
        quien se fue, sin reemplazo o sin fechas) se omiten.
    """
    filas_crudas = _leer_filas_crudas(ruta)
    registros: list[RegistroNovedad] = []
    for fila in filas_crudas:
        registro = _extraer_registro(fila)
        if registro is not None:
            registros.append(registro)
    return registros


def _leer_filas_crudas(ruta: Path | str) -> list[tuple]:
    """Lee las filas de datos crudas (a partir de fila 4) usando openpyxl.

    Se usa lectura posicional porque el archivo de novedades tiene encabezados
    DUPLICADOS (CEDULA y NOMBRE aparecen 2 veces: reemplazo y quien se fue),
    y el dict de ExcelReader los colapsaría.
    """
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = wb[wb.sheetnames[0]]
        filas: list[tuple] = []
        for fila_tuple in hoja.iter_rows(
            min_row=_FILA_ENCABEZADO + 1, values_only=True
        ):
            if all(v is None for v in fila_tuple):
                continue
            filas.append(fila_tuple)
        return filas
    finally:
        wb.close()


def _extraer_registro(fila: tuple) -> RegistroNovedad | None:
    """Convierte una fila cruda en RegistroNovedad; None si está incompleta.

    Los índices son posiciones de la fila (1-based):
        1 = CEDULA reemplazo, 2 = NOMBRE reemplazo, 7 = INICIO, 8 = FIN,
        10 = CEDULA quien se fue, 11 = NOMBRE quien se fue.
    """
    valores = _valores_fila(fila)
    cedula_rep = valores[_COL_CEDULA_REEMPLAZO]
    nombre_rep = valores[_COL_NOMBRE_REEMPLAZO]
    inicio = _parsear_fecha(valores[_COL_INICIO])
    fin = _parsear_fecha(valores[_COL_FIN])
    cedula_fue = valores[_COL_CEDULA_QUIEN_SE_FUE]
    nombre_fue = valores[_COL_NOMBRE_QUIEN_SE_FUE]

    if (
        cedula_fue is None
        or nombre_fue is None
        or cedula_rep is None
        or nombre_rep is None
        or inicio is None
        or fin is None
    ):
        return None
    if fin < inicio:
        return None
    return RegistroNovedad(
        cedula_quien_se_fue=cedula_fue,
        nombre_quien_se_fue=_texto(nombre_fue),
        inicio=inicio,
        fin=fin,
        cedula_reemplazo=cedula_rep,
        nombre_reemplazo=_texto(nombre_rep),
    )


def _valores_fila(fila: tuple) -> dict[int, Any]:
    """Mapa índice (1-based) → valor de la fila cruda (0-based)."""
    valores: dict[int, Any] = {}
    for i, valor in enumerate(fila, start=1):
        valores[i] = valor
    return valores


def _texto(valor) -> str:
    """Convierte a str limpio (o vacío si None)."""
    return str(valor).strip() if valor is not None else ""


def _parsear_fecha(valor) -> date | None:
    """Convierte un valor a date; None si no se puede parsear.

    Acepta date, datetime y strings en formatos dd/mm/yyyy, yyyy-mm-dd,
    dd-mm-yyyy, yyyy/mm/dd.
    """
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if not isinstance(valor, str):
        return None
    texto = valor.strip()
    for fmt in _FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def aplicar_novedades(
    writer: ExcelWriter,
    ruta_novedades: Path | str,
    logger=None,
) -> list[ResultadoNovedad]:
    """Aplica los reemplazos de novedades a la pestaña Ventas.

    Para cada registro: busca en Ventas las filas donde R == cédula de quien
    se fue y F dentro de [inicio, fin], y escribe en R/S la cédula y nombre
    del REEMPLAZO. Devuelve auditoría por persona.

    Args:
        writer: ExcelWriter abierto sobre el archivo de salida.
        ruta_novedades: ruta al archivo de novedades.
        logger: opcional, recibe str para logging.

    Returns:
        Lista de ResultadoNovedad (uno por registro procesado).
    """
    registros = leer_novedades(ruta_novedades)
    if not registros:
        _log(logger, "[novedades] Sin registros válidos en el archivo.")
        return []

    ws = writer._workbook["Ventas"]
    resultados: list[ResultadoNovedad] = []
    for registro in registros:
        resultado = _aplicar_un_registro(writer, ws, registro)
        resultados.append(resultado)
        _log(
            logger,
            f"[novedades] {resultado.nombre_quien_se_fue} ({resultado.cedula_quien_se_fue}) "
            f"→ reemplazado por {resultado.nombre_reemplazo} ({resultado.cedula_reemplazo}): "
            f"{resultado.cambiados}/{resultado.encontrados} cambiados "
            f"[{resultado.intervalo}] {resultado.estado.value}",
        )
    return resultados


def _aplicar_un_registro(
    writer: ExcelWriter,
    ws,
    registro: RegistroNovedad,
) -> ResultadoNovedad:
    """Aplica un único registro de novedades y devuelve su auditoría.

    Busca las filas de Ventas donde R == cédula de quien se fue; de esas,
    las que tienen F dentro de [inicio, fin] se escriben con el reemplazo.
    """
    filas_match: list[int] = []
    for fila in range(2, ws.max_row + 1):
        r_valor = ws.cell(row=fila, column=_COL_VENTAS_R).value
        if r_valor is None:
            continue
        if not _valores_iguales(r_valor, registro.cedula_quien_se_fue):
            continue
        filas_match.append(fila)

    if not filas_match:
        return ResultadoNovedad(
            nombre_quien_se_fue=registro.nombre_quien_se_fue,
            cedula_quien_se_fue=registro.cedula_quien_se_fue,
            nombre_reemplazo=registro.nombre_reemplazo,
            cedula_reemplazo=registro.cedula_reemplazo,
            encontrados=0,
            cambiados=0,
            intervalo="",
            estado=EstadoNovedad.SIN_DATOS,
        )

    # Filtrar por fecha dentro del rango y calcular el intervalo real.
    fechas_en_rango: list[date] = []
    filas_a_marcar: list[int] = []
    for fila in filas_match:
        fecha = _fecha_fila(ws, fila)
        if fecha is None:
            continue
        if registro.inicio <= fecha <= registro.fin:
            filas_a_marcar.append(fila)
            fechas_en_rango.append(fecha)

    # Escribir el REEMPLAZO (A/B) en R/S de las filas que cumplen el rango.
    for fila in filas_a_marcar:
        ws.cell(row=fila, column=_COL_VENTAS_R, value=registro.cedula_reemplazo)
        ws.cell(row=fila, column=_COL_VENTAS_S, value=registro.nombre_reemplazo)

    intervalo = _formatear_intervalo(fechas_en_rango) if fechas_en_rango else ""
    if not fechas_en_rango or len(set(fechas_en_rango)) < _dias_en_rango(registro):
        estado = EstadoNovedad.PARCIAL
    else:
        estado = EstadoNovedad.COMPLETO

    return ResultadoNovedad(
        nombre_quien_se_fue=registro.nombre_quien_se_fue,
        cedula_quien_se_fue=registro.cedula_quien_se_fue,
        nombre_reemplazo=registro.nombre_reemplazo,
        cedula_reemplazo=registro.cedula_reemplazo,
        encontrados=len(filas_match),
        cambiados=len(filas_a_marcar),
        intervalo=intervalo,
        estado=estado,
    )


def _fecha_fila(ws, fila: int) -> date | None:
    """Devuelve la fecha (col F) de una fila de Ventas o None."""
    return _parsear_fecha(ws.cell(row=fila, column=_COL_VENTAS_FECHA).value)


def _valores_iguales(a, b) -> bool:
    """Compara dos valores de celda (int vs str numérico)."""
    if a is None or b is None:
        return False
    try:
        return int(a) == int(b)
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def _dias_en_rango(registro: RegistroNovedad) -> int:
    """Días calendario cubiertos por el rango [inicio, fin] (inclusivo)."""
    return (registro.fin - registro.inicio).days + 1


def _formatear_intervalo(fechas: list[date]) -> str:
    """Formatea el intervalo real de fechas encontradas (dd/mm/aaaa)."""
    if not fechas:
        return ""
    menor = min(fechas).strftime("%d/%m/%Y")
    mayor = max(fechas).strftime("%d/%m/%Y")
    return f"{menor} → {mayor}"


def _log(logger, mensaje: str) -> None:
    """Log opcional hacia el logger del proceso."""
    if logger is not None:
        logger.info(mensaje)
