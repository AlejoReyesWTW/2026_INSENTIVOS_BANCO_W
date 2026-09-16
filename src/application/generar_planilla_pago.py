"""Genera el segundo archivo de salida (Planilla de pago) desde el archivo 1.

Lee la pestaña Ventas del archivo 1 (ya generado con los incentivos
calculados y novedades aplicadas) y crea una hoja con 2 tablas separadas
por 2 columnas en blanco:

    [Tabla CAJEROS]              | 2 cols vacías |  [Tabla SUBGERENTES]
    Nombre Cajero | Cédula | Suma               |  Nombre Subg. | Cédula | Suma
    ...           | ...    | ...                |  ...          | ...    | ...
    TOTAL         |        | Σ                  |  TOTAL        |        | Σ

- Tabla CAJEROS: nombre = K (11), cédula = J (10), suma = SUM(L (12)) por cédula.
- Tabla SUBGERENTES: usa R (18) y S (19) (que ya tienen al reemplazo si hubo
  novedad), suma = SUM(M (13)) por cédula.
- Al final de cada tabla, una fila TOTAL con la sumatoria de la columna suma.

No escribe fórmulas: solo valores (evita la corrupción de Excel).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

# Columnas de Ventas (1-based) en el archivo 1.
_COL_NOMBRE_CAJERO = 11  # K
_COL_CC_CAJERO = 10  # J
_COL_SUMA_CAJERO = 12  # L
_COL_SUMA_SUBGERENTE = 13  # M
_COL_CC_SUBGERENTE = 18  # R (CC_NOVEDAD_SUBGERENTE: reemplazo si hubo)
_COL_NOMBRE_SUBGERENTE = 19  # S

# Columnas de la hoja de salida (1-based).
_TABLA_CAJEROS_INICIO = 1
_TABLA_SUBGERENTES_INICIO = 6  # 3 de tabla1 + 2 en blanco (cols 4,5) → tabla2 col 6

_ENCABEZADOS_CAJEROS = ("Nombre Cajero", "Cédula", "Suma Cajero")
_ENCABEZADOS_SUBGERENTES = ("Nombre Subgerente", "Cédula", "Suma Subgerente")

# Estilos de encabezados: negrita + fondo azul suave.
_ESTILO_ENCABEZADO = Font(bold=True)
_FONDO_ENCABEZADO = PatternFill(
    start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"
)
_ALINEACION_CENTRO = Alignment(horizontal="center")


@dataclass(frozen=True)
class ResumenFuncionario:
    """Resumen agregado por funcionario (una fila de la tabla)."""

    nombre: str
    cedula: object
    suma: float


def agrupar_cajeros(ws_ventas) -> list[ResumenFuncionario]:
    """Agrupa las ventas por cédula de CAJERO y suma la columna L.

    Args:
        ws_ventas: hoja "Ventas" del archivo 1 (openpyxl worksheet).

    Returns:
        Lista de ResumenFuncionario (nombre, cedula, suma) por cajero,
        ordenada de mayor a menor por la suma.
    """
    acumulador: dict = {}
    # iter_rows(values_only=True) es el método rápido y compatible con
    # read_only; acceder con ws.cell() en read_only es MUY lento / no
    # soportado con decenas de miles de filas.
    for fila in ws_ventas.iter_rows(min_row=2, values_only=True):
        cedula = _indice(fila, _COL_CC_CAJERO)
        if cedula is None or str(cedula).strip() == "":
            continue
        nombre = _indice(fila, _COL_NOMBRE_CAJERO)
        suma = _a_numero(_indice(fila, _COL_SUMA_CAJERO))
        clave = _clave_cedula(cedula)
        if clave in acumulador:
            acumulador[clave]["suma"] += suma
        else:
            acumulador[clave] = {
                "nombre": _texto(nombre),
                "cedula": cedula,
                "suma": suma,
            }
    resultados = [
        ResumenFuncionario(nombre=d["nombre"], cedula=d["cedula"], suma=d["suma"])
        for d in acumulador.values()
    ]
    return sorted(resultados, key=lambda resumen: resumen.suma, reverse=True)


def agrupar_subgerentes(ws_ventas) -> list[ResumenFuncionario]:
    """Agrupa las ventas por cédula de SUBGERENTE y suma la columna M.

    Usa las columnas R (18) / S (19): ya contienen al reemplazo si la venta
    fue cubierta por una novedad; si no, son réplica del subgerente normal.

    Args:
        ws_ventas: hoja "Ventas" del archivo 1 (openpyxl worksheet).

    Returns:
        Lista de ResumenFuncionario (nombre, cedula, suma) por subgerente,
        ordenada de mayor a menor por la suma.
    """
    acumulador: dict = {}
    for fila in ws_ventas.iter_rows(min_row=2, values_only=True):
        cedula = _indice(fila, _COL_CC_SUBGERENTE)
        if cedula is None or str(cedula).strip() == "":
            continue
        nombre = _indice(fila, _COL_NOMBRE_SUBGERENTE)
        suma = _a_numero(_indice(fila, _COL_SUMA_SUBGERENTE))
        clave = _clave_cedula(cedula)
        if clave in acumulador:
            acumulador[clave]["suma"] += suma
        else:
            acumulador[clave] = {
                "nombre": _texto(nombre),
                "cedula": cedula,
                "suma": suma,
            }
    resultados = [
        ResumenFuncionario(nombre=d["nombre"], cedula=d["cedula"], suma=d["suma"])
        for d in acumulador.values()
    ]
    return sorted(resultados, key=lambda resumen: resumen.suma, reverse=True)


def _indice(fila: tuple, col: int):
    """Devuelve el valor de una columna (1-based) en la tupla de la fila."""
    if len(fila) < col:
        return None
    return fila[col - 1]


def generar_planilla_pago(
    ruta_archivo1: Path | str,
    ruta_destino: Path | str,
) -> Path:
    """Crea el archivo 2 (Planilla de pago) desde el archivo 1.

    Args:
        ruta_archivo1: ruta al archivo 1 ya generado (pestaña Ventas).
        ruta_destino: ruta donde guardar el archivo 2 (.xlsx).

    Returns:
        La ruta de destino (Path).
    """
    wb1 = openpyxl.load_workbook(ruta_archivo1, data_only=True, read_only=True)
    try:
        ws_ventas = wb1["Ventas"]
        cajeros = agrupar_cajeros(ws_ventas)
        subgerentes = agrupar_subgerentes(ws_ventas)
    finally:
        wb1.close()

    # Construir el archivo 2.
    wb2 = openpyxl.Workbook()
    hoja = wb2.active
    if hoja is not None:
        hoja.title = "Planilla de pago"
        _escribir_encabezados(hoja)
        _escribir_tabla(
            hoja,
            inicio_columna=_TABLA_CAJEROS_INICIO,
            fila_inicio=2,
            datos=cajeros,
        )
        _escribir_tabla(
            hoja,
            inicio_columna=_TABLA_SUBGERENTES_INICIO,
            fila_inicio=2,
            datos=subgerentes,
        )
        _escribir_total_tabla(
            hoja,
            inicio_columna=_TABLA_CAJEROS_INICIO,
            fila=2 + len(cajeros),
            total=sum(r.suma for r in cajeros),
        )
        _escribir_total_tabla(
            hoja,
            inicio_columna=_TABLA_SUBGERENTES_INICIO,
            fila=2 + len(subgerentes),
            total=sum(r.suma for r in subgerentes),
        )
        _autoajustar_columnas(hoja)
    wb2.save(ruta_destino)
    wb2.close()
    return Path(ruta_destino)


def _escribir_encabezados(hoja) -> None:
    """Escribe los encabezados de ambas tablas en la fila 1 (bold + azul)."""
    for i, texto in enumerate(_ENCABEZADOS_CAJEROS, start=_TABLA_CAJEROS_INICIO):
        _estilizar_encabezado(hoja.cell(row=1, column=i, value=texto))
    for i, texto in enumerate(
        _ENCABEZADOS_SUBGERENTES, start=_TABLA_SUBGERENTES_INICIO
    ):
        _estilizar_encabezado(hoja.cell(row=1, column=i, value=texto))


def _estilizar_encabezado(celda) -> None:
    """Aplica negrita + fondo azul suave + centrado a una celda de encabezado."""
    celda.font = _ESTILO_ENCABEZADO
    celda.fill = _FONDO_ENCABEZADO
    celda.alignment = _ALINEACION_CENTRO


def _escribir_tabla(hoja, inicio_columna: int, fila_inicio: int, datos: list) -> None:
    """Escribe una tabla (nombre | cédula | suma) empezando en la columna dada."""
    for offset, resumen in enumerate(datos):
        fila = fila_inicio + offset
        hoja.cell(row=fila, column=inicio_columna, value=resumen.nombre)
        hoja.cell(row=fila, column=inicio_columna + 1, value=resumen.cedula)
        hoja.cell(row=fila, column=inicio_columna + 2, value=resumen.suma)


def _escribir_total_tabla(hoja, inicio_columna: int, fila: int, total: float) -> None:
    """Escribe la fila TOTAL al final de una tabla (nombre + suma)."""
    celda_nombre = hoja.cell(row=fila, column=inicio_columna, value="TOTAL")
    celda_suma = hoja.cell(row=fila, column=inicio_columna + 2, value=total)
    for celda in (celda_nombre, celda_suma):
        celda.font = _ESTILO_ENCABEZADO  # negrita para resaltar el total


def _autoajustar_columnas(hoja) -> None:
    """Ajusta el ancho de cada columna según su contenido (máx. 45)."""
    for col_cells in hoja.columns:
        col_letter = col_cells[0].column_letter
        max_largo = 0
        for celda in col_cells:
            if celda.value is None:
                continue
            largo = len(str(celda.value))
            if largo > max_largo:
                max_largo = largo
        ancho = min(max_largo + 2, 45)
        hoja.column_dimensions[col_letter].width = max(ancho, 8)


def _a_numero(valor) -> float:
    """Convierte un valor a float; 0 si no es numérico."""
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        try:
            return float(valor)
        except (TypeError, ValueError, OverflowError):
            return 0.0
    try:
        return float(str(valor).strip())
    except (TypeError, ValueError):
        return 0.0


def _clave_cedula(cedula) -> str:
    """Normaliza una cédula para usarla como clave de agrupación."""
    try:
        return str(int(cedula))
    except (TypeError, ValueError):
        return str(cedula).strip()


def _texto(valor) -> str:
    """Convierte a str limpio (o vacío si None)."""
    return str(valor).strip() if valor is not None else ""
